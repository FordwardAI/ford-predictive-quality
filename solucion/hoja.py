"""Hoja de códigos prioritarios (P8, E3): planilla CSV/XLSX e imprimible de una página.

    .venv/bin/python -m solucion.hoja --csv "<Dataset QLS Inspección Adicional.csv>" \\
        --catalogo "<Códigos de catálogo.csv>" --dia 190 --salida DIR [--programa programa.csv] [--cupo N]

Entradas: la tabla por VIN, el programa del día (CSV con columnas `unidad,codigo`) y el cupo.
Sin programa (demo) se simula con los VIN del Día del VIN pedido, reemplazados por ids
ficticios U-0001…; sin cupo, se usa el cupo diario del 5 %. La hoja prioriza códigos, nunca
unidades, y nunca contiene un VIN. Tiene tasas por código: se escribe solo en `--salida`, fuera del repo.
"""
import argparse
import collections
import csv
import dataclasses
import html
import json
import re
import sys
import zipfile
from pathlib import Path

import numpy as np

import importlib.util

from . import cupo as cupo_mod
from . import referencias  # noqa: F401  Registra las familias que reconstruye eleccion.ganadora.
from . import datos
from .cupo import SEMILLA_DESEMPATE, fuente_completa, version_codigo
from .datos import MARGEN, TRAMOS
from .eleccion import ganadora
from .puntaje import Contexto, semilla_dia, wilson

from catalog_groups import chi_square  # noqa: E402  (research/ queda en sys.path al importar datos)

if importlib.util.find_spec("solucion.ml"):  # P4, si ya existe: la ganadora puede ser un modelo.
    importlib.import_module("solucion.ml")

RESULTADOS = datos.RAIZ / "solucion" / "resultados"
DIA_DEMO = 190
P_PROVISORIO = 20  # Mínimo por código: 1 auditoría cada P días, hasta que P5 elija P en validación.
VENTANA_SIN_PROPIA = 60  # n que se muestra para un predictor sin ventana propia (decaimiento, ML).
SEMILLA_PROGRAMA = 20261008
VALIDACION = TRAMOS["validacion"]
COLUMNAS = ["Código", "Cantidad sugerida", "Tasa reciente (rango 95 %)", "n (auditados con resultado)",
            "Veces la tasa general", "Unidades programadas", "Acumulado", "Mercado de destino",
            "Versión (dominante)", "Motor (dominante)", "Tracción"]
LIMITES = [
    "Es una base ficticia y reúne solo auditados con actividad QLS (el sistema donde la verificación de "
    "calidad registra incidencias y reparaciones).",
    "El tramo de la prueba final ya se había explorado y la tasa por mercado en ese tramo ya se vio: la cifra "
    "final puede ser optimista.",
    "Que los auditados se eligieron al azar es un supuesto de Ford.",
    "En planta solo se conocería el resultado de las unidades que se auditan.",
    "El Día del VIN aproxima el día de la auditoría.",
    "No indica impacto ni ahorro en planta, menos calibraciones, validez para unidades no auditadas ni causas.",
]


# --- Formato -------------------------------------------------------------------------------------------------

def _pct(x):
    return "—" if x is None else f"{100 * x:.1f} %".replace(".", ",")


def _entero(n):
    return f"{int(round(n)):,}".replace(",", ".")


def _decimal(x, cifras=2):
    return "—" if x is None else f"{x:.{cifras}f}".replace(".", ",")


def _dias(desde, hasta):
    return f"días hasta el {hasta}" if not desde else f"días {desde}–{hasta}"


# --- Programa del día ----------------------------------------------------------------------------------------

def leer_programa(path):
    """Programa del día: lista de (unidad, codigo) en el orden del archivo."""
    with Path(path).open(encoding="utf-8-sig", newline="") as fuente:
        filas = [(r["unidad"].strip(), r["codigo"].strip()) for r in csv.DictReader(fuente)]
    assert filas and all(u and c for u, c in filas), "El programa necesita columnas unidad,codigo sin vacíos"
    assert len({u for u, _ in filas}) == len(filas), "Unidad repetida en el programa"
    return filas


def simular_programa(tabla, t, semilla=SEMILLA_PROGRAMA):
    """Demo: las unidades del Día del VIN t con ids ficticios. El VIN no sale de esta función.

    Se ordenan por código y se barajan con semilla: el id solo indica la posición en ese orden.
    """
    codigos = sorted(v.codigo for v in tabla.por_dia(t, t).get(t, []))
    orden = np.random.default_rng([semilla, t]).permutation(len(codigos))
    return [(f"U-{i + 1:04d}", codigos[j]) for i, j in enumerate(orden)]


def escribir_programa(programa, destino):
    with Path(destino).open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["unidad", "codigo"])
        w.writerows(programa)


# --- Cantidad sugerida y traspaso ----------------------------------------------------------------------------

def minimo_p():
    """P del mínimo por código: el elegido por P5 si existe, si no el provisorio."""
    archivo = RESULTADOS / "p5.json"
    if archivo.exists():
        p = json.loads(archivo.read_text()).get("minimo_por_codigo", {}).get("P")
        if p:
            return int(p), "elegido en validación (P5)"
    return P_PROVISORIO, "provisorio, hasta que P5 lo elija en validación"


def sugerir(ranking, programadas, vencidos, cupo_dia):
    """Cantidad sugerida por código.

    Primero el mínimo por código: una unidad a cada código vencido, en el orden recibido (el más antiguo
    primero) y sin pasar el cupo. Después se llena el resto bajando por el ranking.
    Devuelve (por_ranking {código: cantidad}, exploracion [códigos], sin_cubrir).
    """
    exploracion = [c for c in vencidos if programadas.get(c, 0) > 0][:cupo_dia]
    ya = collections.Counter(exploracion)
    pendiente = cupo_dia - len(exploracion)
    por_ranking = {}
    for c in ranking:
        q = min(programadas[c] - ya[c], pendiente)
        if q > 0:
            por_ranking[c] = q
            pendiente -= q
    return por_ranking, exploracion, pendiente


def reasignar(sugeridas, ranking, llegados):
    """Si un código no llega a la playa (o llega con menos unidades), lo pendiente baja por el ranking.

    `sugeridas` {código: cantidad total}, `ranking` [códigos, de mayor a menor tasa], `llegados`
    {código: unidades en la playa}. Lo pendiente de cada código pasa a los códigos siguientes del
    ranking que tengan unidades sin asignar. Si no alcanza, a los de más arriba que hayan llegado con más
    unidades que las programadas. Solo lo que no entra al agotarse el ranking va al azar.
    Devuelve ({código: cantidad}, cantidad a completar al azar).
    """
    assert set(sugeridas) <= set(ranking), "Un código sugerido no está en el ranking"
    nueva = {c: min(q, llegados.get(c, 0)) for c, q in sugeridas.items()}
    faltante = {c: q - nueva[c] for c, q in sugeridas.items()}
    azar = 0
    for i, c in enumerate(ranking):
        pendiente = faltante.get(c, 0)
        for d in ranking[i + 1:] + ranking[:i]:
            if not pendiente:
                break
            mover = min(llegados.get(d, 0) - nueva.get(d, 0), pendiente)
            if mover > 0:
                nueva[d] = nueva.get(d, 0) + mover
                pendiente -= mover
        azar += pendiente
    return {c: q for c, q in nueva.items() if q > 0}, azar


# --- Armado --------------------------------------------------------------------------------------------------

@dataclasses.dataclass
class Fila:
    codigo: str
    sugerida: int
    tasa: float
    rango: tuple
    n: int
    cal: int
    veces: float | None
    programadas: int
    acumulado: int
    mercado: str
    version: str
    motor: str
    traccion: str


@dataclasses.dataclass
class Hoja:
    dia: int
    cupo: int
    cupo_5: int  # El cupo diario del 5 % del programa: único corte con cifra evaluada.
    programadas: int
    filas: list  # Ranking completo de los códigos programados.
    exploracion: list  # [(código, último día con resultado o None)].
    unidades: list  # [(código, motivo, [ids])].
    sin_cubrir: int
    predictor: str
    suavizada: bool
    ventana: tuple
    general: float
    n_ventana: int
    minimo: tuple  # (P, origen).
    evaluacion: dict | None
    por_que: list
    agrupaciones: dict
    ventana_mercado: tuple
    programa_simulado: bool


def ventana(predictor, t):
    """(desde, hasta) de los resultados que usa el predictor; 60 días si no tiene ventana propia."""
    hasta = t - MARGEN
    if predictor.familia == "tasa_fija":
        return None, predictor.hasta
    v = getattr(predictor, "ventana", None) or VENTANA_SIN_PROPIA
    return hasta - v + 1, hasta


def _atributos(tabla, codigo):
    a = tabla.catalogo.get(codigo, {})
    falta = "sin dato en el catálogo"
    return (a.get("mercado", falta), a.get("traccion_dominante", falta), a.get("motor_dominante", falta),
            a.get("traccion", falta))


def _ultimos(conocidas):
    """Último Día con resultado conocido de cada código."""
    salida = {}
    for i, c in enumerate(conocidas.codigos):
        nz = np.nonzero(conocidas.n[i])[0]
        if len(nz):
            salida[c] = int(conocidas.dias[nz[-1]])
    return salida


def agrupaciones(tabla):
    """χ² de cada agrupación con las etiquetas de validación 155–194 (bloque explicativo, no predictor)."""
    claves = {"mercado de destino": "mercado", "motor": "motor_dominante", "tracción": "traccion",
              "versión": "traccion_dominante"}
    salida = {}
    vins = tabla.tramo("validacion")
    for nombre, clave in claves.items():
        grupos = collections.defaultdict(lambda: [0, 0])
        for v in vins:
            g = tabla.catalogo.get(v.codigo, {}).get(clave)
            if g is not None:
                grupos[g][0] += 1
                grupos[g][1] += v.calibrada
        x = chi_square({g: tuple(ab) for g, ab in grupos.items()})
        salida[nombre] = {"chi2": x, "grupos": len(grupos)}
    return salida


def evaluacion_ganadora():
    """Cifras de validación de la ganadora (eleccion.json); None si todavía no hay elección."""
    archivo = RESULTADOS / "eleccion.json"
    if not archivo.exists():
        return None
    g = json.loads(archivo.read_text())["ganadora"]
    azar = g["precision_cupo"] / g["veces_azar"] if g.get("veces_azar") else None
    return {"alternativa": g["alternativa"], "precision": g["precision_cupo"], "rango": g["precision_rango95"],
            "azar": azar, "lectura": g["lectura"], "calificador": g["calificador"]}


def armar(tabla, programa, t, cupo_dia=None, predictor=None, minimo=None, evaluacion="eleccion",
          programa_simulado=False):
    """Arma la hoja del día t. Las tasas usan solo resultados de Día <= t−5."""
    predictor = predictor or ganadora(tabla)
    assert not getattr(predictor, "por_vin", False), "La hoja prioriza códigos, no unidades"
    minimo = minimo or minimo_p()
    evaluacion = evaluacion_ganadora() if evaluacion == "eleccion" else evaluacion
    programadas = collections.Counter(c for _, c in programa)
    cupo_5 = cupo_mod.cupo(len(programa))
    cupo_dia = cupo_5 if cupo_dia is None else cupo_dia
    fuente = fuente_completa(tabla)
    ctx = Contexto(fuente, t)
    codigos = sorted(programadas)
    tasas = predictor.puntuar(ctx, codigos)
    desde, hasta = ventana(predictor, t)
    conocidas = ctx.conocidas(hasta, desde=desde)
    crudos, general = conocidas.por_codigo(), conocidas.general()
    n_ventana = int(conocidas.n.sum())

    rng = semilla_dia(SEMILLA_DESEMPATE, t)
    desempate = dict(zip(codigos, rng.random(len(codigos))))
    ranking = sorted(codigos, key=lambda c: (-tasas[c], desempate[c]))

    ultimos = _ultimos(ctx.conocidas(t - MARGEN))
    limite = t - MARGEN - minimo[0] + 1  # Sin resultado conocido en los últimos P días con resultado.
    posicion = {c: i for i, c in enumerate(ranking)}
    vencidos = sorted((c for c in codigos if ultimos.get(c, -1) < limite),
                      key=lambda c: (ultimos.get(c, -1), posicion[c]))
    por_ranking, exploracion, sin_cubrir = sugerir(ranking, programadas, vencidos, cupo_dia)

    filas, acumulado = [], 0
    for c in ranking:
        n, cal = crudos.get(c, (0.0, 0.0))
        acumulado += programadas[c]
        filas.append(Fila(c, por_ranking.get(c, 0), tasas[c], wilson(cal, n) if n else None, int(n), int(cal),
                          tasas[c] / general if general else None, programadas[c], acumulado,
                          *_atributos(tabla, c)))

    orden = collections.defaultdict(list)
    for u, c in programa:
        orden[c].append(u)
    unidades = [(c, "Mínimo por código", orden[c][:1]) for c in exploracion]
    unidades += [(c, "Prioridad", orden[c][c in exploracion:(c in exploracion) + q]) for c, q in por_ranking.items()]

    # «Por qué este código»: mercado de destino con resultados de validación hasta t−5.
    hasta_m = t - MARGEN
    desde_m = VALIDACION[0] if hasta_m >= VALIDACION[0] else max(0, hasta_m - VENTANA_SIN_PROPIA + 1)
    conocidas_m = ctx.conocidas(hasta_m, desde=desde_m)
    general_m = conocidas_m.general()
    por_mercado = collections.defaultdict(lambda: [0.0, 0.0])
    for c, (n, cal) in conocidas_m.por_codigo().items():
        m = tabla.mercado(c)
        if m is not None:
            por_mercado[m][0] += n
            por_mercado[m][1] += cal
    por_que = []
    for f in filas:
        if not f.sugerida:
            continue
        n_m, cal_m = por_mercado.get(f.mercado, (0.0, 0.0))
        por_que.append({"codigo": f.codigo, "tasa": f.tasa, "n": f.n, "cal": f.cal, "veces": f.veces,
                        "mercado": f.mercado, "tasa_mercado": cal_m / n_m if n_m else None, "n_mercado": int(n_m),
                        "general_mercado": general_m})

    return Hoja(dia=t, cupo=cupo_dia, cupo_5=cupo_5, programadas=len(programa), filas=filas,
                exploracion=[(c, ultimos.get(c)) for c in exploracion], unidades=unidades, sin_cubrir=sin_cubrir,
                predictor=predictor.nombre,
                suavizada=not (predictor.familia in ("tasa_fija", "movil") and not predictor.parametros.get("peso")),
                ventana=(desde, hasta), general=general, n_ventana=n_ventana, minimo=minimo, evaluacion=evaluacion,
                por_que=por_que, agrupaciones=agrupaciones(tabla), ventana_mercado=(desde_m, hasta_m),
                programa_simulado=programa_simulado)


# --- Textos --------------------------------------------------------------------------------------------------

def _rango_txt(f):
    if f.rango is None:
        return f"{_pct(f.tasa)} (sin resultados: tasa general)"
    return f"{_pct(f.tasa)} ({_decimal(100 * f.rango[0], 1)} a {_pct(f.rango[1])})"


def _valores(f):
    return [f.codigo, f.sugerida, _rango_txt(f), f.n, _decimal(f.veces), f.programadas, f.acumulado,
            f.mercado, f.version, f.motor, f.traccion]


def _lectura(x):
    return {"mejora": "mejor que el azar", "inconcluso": "no se distingue del azar",
            "peor": "peor que el azar"}.get(x, x)


def textos(h):
    """Textos comunes a las tres salidas."""
    lo, hi = h.ventana
    e = h.evaluacion
    if e is None:
        evaluacion = ("Sin cifra evaluada todavía: el orden usa una tasa provisoria hasta que se elija el método "
                      "en validación.")
    else:
        calificador = re.sub(r"\d{4,}", lambda m: _entero(int(m.group())), e["calificador"])
        evaluacion = (f"Con el cupo del 5 %, {calificador}: de cada 100 unidades elegidas se calibrarían "
                      f"{_decimal(100 * e['precision'], 1)} (rango del 95 %: {_decimal(100 * e['rango'][0], 1)} a "
                      f"{_decimal(100 * e['rango'][1], 1)}), contra {_decimal(100 * e['azar'], 1)} al azar. "
                      f"Resultado: {_lectura(e['lectura'])}. La prueba final todavía no se corrió.")
    corte = (f"El cupo del 5 % ({h.cupo_5} de {h.programadas} unidades) es el único corte con cifra evaluada; "
             "por debajo del cupo, sin cifra evaluada.")
    tasa = (f"Tasa reciente: proporción calibrada del código entre sus auditados con resultado conocido "
            f"({_dias(lo, hi)}). Tasa general: {_pct(h.general)} (entre auditados con actividad QLS, "
            f"{_dias(lo, hi)}, base ficticia, n = {_entero(h.n_ventana)} VIN).")
    if h.suavizada:
        tasa += (" La tasa que ordena se ajusta hacia la tasa general (o la de su mercado) para que un código con "
                 "pocos resultados no suba por casualidad; el rango y el n son los del código.")
    minimo = (f"Mínimo por código: una unidad a cada código programado sin resultado conocido en los últimos "
              f"{h.minimo[0]} días ({h.minimo[1]}), el más antiguo primero. Mantiene al día la tasa de los códigos "
              "poco elegidos.")
    uso = ["Al inicio del día, repartir el cupo según la cantidad sugerida por código.",
           "En cada ronda, tomar en la playa de despacho unidades de los códigos con cantidad pendiente: el código "
           "está en la etiqueta del parabrisas. Dentro de un código, cualquier unidad sirve.",
           "Si un código no llega (o llega con menos unidades), lo pendiente pasa a los códigos siguientes de la "
           "lista que sí llegaron. Solo si se agota la lista, completar al azar."]
    return {"evaluacion": evaluacion, "corte": corte, "tasa": tasa, "minimo": minimo, "uso": uso}


def _por_que_txt(h, p):
    lo, hi = h.ventana_mercado
    codigo = (f"{p['codigo']}: {_pct(p['tasa'])}, {_decimal(p['veces'])} veces la tasa general"
              + (f" ({p['cal']} calibradas de {p['n']})." if p["n"] else " (sin resultados propios)."))
    if p["tasa_mercado"] is None:
        return codigo + f" Mercado de destino {p['mercado']}: sin resultados en {_dias(lo, hi)}."
    return codigo + (f" Su mercado de destino, {p['mercado']}, se calibró en {_pct(p['tasa_mercado'])} "
                     f"(n = {_entero(p['n_mercado'])}) contra {_pct(p['general_mercado'])} en general, {_dias(lo, hi)}.")


def _agrupaciones_txt(h):
    partes = [f"{k} {_decimal(v['chi2'], 1)}" for k, v in h.agrupaciones.items() if v["chi2"] is not None]
    return (f"Validación {VALIDACION[0]}–{VALIDACION[1]}: el mercado de destino es la agrupación del catálogo que más "
            f"separa las tasas (diferencia entre grupos, χ²: {' · '.join(partes)}; más alto, más diferencia). "
            "Es una asociación, no una causa.")


PENDIENTES = ["Dónde mirar (componente más probable): pendiente de prueba final. Entra solo si se sostiene.",
              "Códigos que casi no se calibran: pendiente de prueba final. Entran solo si se sostienen."]


# --- Salidas -------------------------------------------------------------------------------------------------

def escribir_csv(h, destino):
    with Path(destino).open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(COLUMNAS + ["Fila"])
        for fila in h.filas:
            w.writerow(_valores(fila) + ["Prioridad"])
        por_codigo = {f.codigo: f for f in h.filas}
        for c, _ in h.exploracion:
            w.writerow(_valores(dataclasses.replace(por_codigo[c], sugerida=1)) + ["Mínimo por código"])


def escribir_xlsx(h, destino):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    tx = textos(h)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Hoja del día"
    negrita, azul = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    gris, claro = Font(color="6B7280", italic=True), PatternFill("solid", fgColor="EAF1FB")
    borde = Border(bottom=Side(style="thin", color="C7CED9"))
    envolver = Alignment(wrap_text=True, vertical="top")
    ancho = len(COLUMNAS)

    def texto(fila, valor, fuente=None, alto=None):
        hoja.cell(fila, 1, valor).alignment = envolver
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
        if fuente:
            hoja.cell(fila, 1).font = fuente
        if alto:
            hoja.row_dimensions[fila].height = alto

    texto(1, f"Hoja de códigos prioritarios · Día {h.dia}", Font(bold=True, size=14))
    texto(2, f"{_entero(h.programadas)} unidades programadas · cupo del día: {h.cupo} · {len(h.filas)} códigos",
          Font(bold=True))
    texto(3, tx["evaluacion"], None, 45)
    texto(4, tx["corte"] + " " + tx["tasa"], Font(size=9), 45)
    cabecera = 6
    for j, nombre in enumerate(COLUMNAS, 1):
        celda = hoja.cell(cabecera, j, nombre)
        celda.font, celda.fill = negrita, azul
        celda.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    hoja.row_dimensions[cabecera].height = 32
    fila = cabecera + 1
    for f in h.filas:
        for j, valor in enumerate(_valores(f), 1):
            celda = hoja.cell(fila, j, f.veces if j == 5 and f.veces is not None else valor)
            celda.border = borde
            if j == 5:
                celda.number_format = "0.00"
            if f.sugerida:
                celda.fill = claro
            elif j > 1:
                celda.font = Font(color="6B7280")
        hoja.cell(fila, 1).font = Font(bold=True)
        fila += 1
    texto(fila, "Las filas resaltadas llenan el cupo. Las demás completan el ranking: sin cifra evaluada.", gris)
    fila += 2
    texto(fila, "Mínimo por código (filas aparte)", Font(bold=True))
    fila += 1
    texto(fila, tx["minimo"], Font(size=9), 30)
    fila += 1
    if h.exploracion:
        for c, ultimo in h.exploracion:
            texto(fila, f"{c}: 1 unidad · último resultado conocido: "
                  + (f"Día {ultimo}" if ultimo is not None else "ninguno"))
            fila += 1
    else:
        texto(fila, "Hoy ningún código programado está vencido.", gris)
        fila += 1
    fila += 1
    texto(fila, "Por qué este código", Font(bold=True))
    fila += 1
    for p in h.por_que:
        texto(fila, _por_que_txt(h, p), None, 30)
        fila += 1
    texto(fila, _agrupaciones_txt(h), Font(size=9), 30)
    fila += 2
    for linea in PENDIENTES:
        texto(fila, linea, gris)
        fila += 1
    fila += 1
    texto(fila, "Límites", Font(bold=True))
    fila += 1
    for linea in LIMITES:
        texto(fila, "• " + linea, Font(size=9))
        fila += 1
    for j, w in enumerate([10, 11, 30, 13, 11, 12, 11, 16, 14, 16, 10], 1):
        hoja.column_dimensions[get_column_letter(j)].width = w
    hoja.freeze_panes = hoja.cell(cabecera + 1, 1)
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.sheet_properties.pageSetUpPr.fitToPage = True

    lista = libro.create_sheet("Unidades sugeridas")
    for j, nombre in enumerate(["Código", "Motivo", "Unidad"], 1):
        celda = lista.cell(1, j, nombre)
        celda.font, celda.fill = negrita, azul
    fila = 2
    for c, motivo, ids in h.unidades:
        for u in ids:
            for j, valor in enumerate([c, motivo, u], 1):
                lista.cell(fila, j, valor).border = borde
            fila += 1
    for letra, w in zip("ABC", (12, 20, 12)):
        lista.column_dimensions[letra].width = w
    lista.freeze_panes = "A2"
    libro.save(destino)


ESTILO = """
:root{--tinta:#1c2230;--suave:#5b6475;--linea:#d6dbe3;--acento:#1f4e79;--acento-claro:#eaf1fb;--fondo:#fff;
--aviso:#7a4b00;--aviso-claro:#fff4dc}
*{box-sizing:border-box}html{background:var(--fondo)}
body{margin:0 auto;max-width:1180px;padding:16px;color:var(--tinta);background:var(--fondo);
font:12.5px/1.35 system-ui,-apple-system,"Segoe UI",sans-serif}
h1{font-size:19px;margin:0}h2{font-size:12.5px;margin:10px 0 4px;color:var(--acento);text-transform:uppercase;
letter-spacing:.03em}p{margin:3px 0}ul,ol{margin:3px 0;padding-left:16px}li{margin:1px 0}
.cabeza{display:flex;justify-content:space-between;align-items:baseline;gap:12px;border-bottom:2px solid var(--acento);
padding-bottom:4px}.suave{color:var(--suave)}.chico{font-size:11px}
.caja{border-left:4px solid var(--acento);background:var(--acento-claro);padding:6px 9px;margin:8px 0;border-radius:3px}
.grilla{display:grid;grid-template-columns:minmax(0,2.3fr) minmax(0,1fr);gap:14px}td{white-space:nowrap}
table{width:100%;border-collapse:collapse}th,td{padding:2px 5px;border-bottom:1px solid var(--linea);text-align:left;
vertical-align:top}th{background:var(--acento);color:#fff;font-size:10.5px;font-weight:600;line-height:1.15}
td.num,th.num{text-align:right;font-variant-numeric:tabular-nums}tr.cupo td{background:var(--acento-claro)}
tr.resto td{color:var(--suave)}tr.linea td{border-bottom:2px solid var(--acento);color:var(--acento);font-size:10.5px;
background:none}.pendiente{background:var(--aviso-claro);color:var(--aviso);padding:4px 8px;border-radius:3px;margin:4px 0}
.cod{font-weight:700;letter-spacing:.02em}.unid{font-variant-numeric:tabular-nums}
@media screen and (max-width:760px){.grilla{grid-template-columns:1fr}table{font-size:11px}td{white-space:normal}}
@page{size:A4 landscape;margin:7mm}
@media print{body{max-width:none;padding:0;font-size:9.6px;line-height:1.25}h1{font-size:14px}
h2{font-size:9.4px;margin:5px 0 2px}.chico{font-size:8.4px}th{font-size:8.2px}td,th{padding:1px 3px}
.caja{margin:4px 0;padding:3px 6px}.grilla{gap:9px}tr.linea td{font-size:7.8px}
*{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
"""


def escribir_html(h, destino):
    tx, e = textos(h), html.escape
    filas, marcado = [], False
    for i, f in enumerate(h.filas, 1):
        clase = "cupo" if f.sugerida else "resto"
        celdas = _valores(f)
        filas.append(f"<tr class='{clase}'><td class='num'>{i}</td><td class='cod'>{e(f.codigo)}</td>"
                     f"<td class='num'><b>{f.sugerida or ''}</b></td><td class='num'>{e(celdas[2])}</td>"
                     f"<td class='num'>{f.n}</td><td class='num'>{e(celdas[4])}</td><td class='num'>{f.programadas}</td>"
                     f"<td class='num'>{f.acumulado}</td><td>{e(f.mercado)}</td><td>{e(f.version)}</td>"
                     f"<td>{e(f.motor)}</td><td>{e(f.traccion)}</td></tr>")
        ultimo_con_cantidad = not any(g.sugerida for g in h.filas[i:])
        if f.sugerida and ultimo_con_cantidad and not marcado:
            marcado = True
            extra = f", con {len(h.exploracion)} del mínimo por código" if h.exploracion else ""
            filas.append(f"<tr class='linea'><td colspan='12'>▲ Hasta acá se llena el cupo del día ({h.cupo}{extra})."
                         " Más abajo: sin cifra evaluada.</td></tr>")
    cabecera = "".join(f"<th{' class=num' if 0 < j < 7 else ''}>{e(c)}</th>" for j, c in enumerate(COLUMNAS))
    exploracion = ("".join(f"<li><span class='cod'>{e(c)}</span>: 1 unidad · último resultado conocido: "
                           f"{'Día ' + str(u) if u is not None else 'ninguno'}</li>" for c, u in h.exploracion)
                   or "<li class='suave'>Hoy ningún código programado está vencido.</li>")
    unidades = "".join(f"<li><span class='cod'>{e(c)}</span> <span class='suave'>({e(m.lower())})</span>: "
                       f"<span class='unid'>{e(', '.join(ids))}</span></li>" for c, m, ids in h.unidades)
    faltan = (f"<p class='pendiente'>El programa no alcanza para el cupo: faltan {h.sin_cubrir}; completar al azar."
              "</p>" if h.sin_cubrir else "")
    programa = ("Programa simulado con las unidades del día en la base (ids ficticios). "
                if h.programa_simulado else "")
    pagina = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Hoja del día {h.dia}</title><style>{ESTILO}</style></head><body>
<div class="cabeza"><h1>Hoja de códigos prioritarios · Día {h.dia}</h1>
<div><b>{_entero(h.programadas)}</b> unidades programadas · cupo del día: <b>{h.cupo}</b> · {len(h.filas)} códigos</div></div>
<div class="caja"><b>Evaluación.</b> {e(tx['evaluacion'])}<div class="chico">{e(tx['corte'])}</div></div>
<div class="grilla"><div>
<table><thead><tr><th class="num">#</th>{cabecera}</tr></thead><tbody>{''.join(filas)}</tbody></table>
<p class="chico suave">{e(programa)}{e(tx['tasa'])} Prioriza códigos: dentro de un código las unidades son
equivalentes. No es la probabilidad de una unidad.</p>
</div><div>
<h2>Cómo se usa</h2><ol class="chico">{''.join(f'<li>{e(x)}</li>' for x in tx['uso'])}</ol>
<h2>Unidades sugeridas</h2>{faltan}<ul class="chico">{unidades}</ul>
<h2>Mínimo por código (filas aparte)</h2><p class="chico">{e(tx['minimo'])}</p><ul class="chico">{exploracion}</ul>
<h2>Por qué este código</h2><ul class="chico">{''.join(f'<li>{e(_por_que_txt(h, p))}</li>' for p in h.por_que)}</ul>
<p class="chico suave">{e(_agrupaciones_txt(h))}</p>
{''.join(f'<p class="pendiente chico">{e(x)}</p>' for x in PENDIENTES)}
<h2>Límites</h2><ul class="chico suave">{''.join(f'<li>{e(x)}</li>' for x in LIMITES)}</ul>
</div></div></body></html>
"""
    Path(destino).write_text(pagina, encoding="utf-8")


def escribir(h, carpeta, programa=None):
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    base = f"hoja-dia-{h.dia}"
    archivos = [carpeta / f"{base}.csv", carpeta / f"{base}.xlsx", carpeta / f"{base}.html"]
    escribir_csv(h, archivos[0])
    escribir_xlsx(h, archivos[1])
    escribir_html(h, archivos[2])
    if programa is not None:
        archivos.append(carpeta / f"programa-dia-{h.dia}.csv")
        escribir_programa(programa, archivos[-1])
    return archivos


def vins_en(archivos, vins):
    """VIN de `vins` que aparecen en los archivos (texto, o el XML interno de un .xlsx)."""
    vins, hallados = set(vins), set()
    for a in archivos:
        a = Path(a)
        if a.suffix == ".xlsx":
            with zipfile.ZipFile(a) as z:
                texto = " ".join(z.read(n).decode("utf-8", "replace") for n in z.namelist())
        else:
            texto = a.read_text(encoding="utf-8-sig")
        hallados |= vins & set(re.findall(r"[A-Za-z0-9]+", texto))
    return hallados


# --- Entrada -------------------------------------------------------------------------------------------------

def construir_hoja(tabla, dia, carpeta, programa_path=None, cupo_dia=None, **kw):
    """Arma y escribe la hoja; verifica que ningún VIN de la tabla aparezca en las salidas."""
    simulado = programa_path is None
    programa = simular_programa(tabla, dia) if simulado else leer_programa(programa_path)
    assert programa, f"No hay unidades programadas para el Día {dia}"
    h = armar(tabla, programa, dia, cupo_dia, programa_simulado=simulado, **kw)
    archivos = escribir(h, carpeta, programa if simulado else None)
    todos = [v.vin for v in tabla.vins] + [v.vin for v in tabla.cohorte]
    assert not vins_en(archivos, todos), "Una salida de la hoja contiene un VIN"
    return h, archivos


def resumen(h, archivos, tabla):
    """Solo conteos: sin tasas por código ni VIN."""
    return {"pieza": "P8 hoja de códigos prioritarios", "dia": h.dia,
            "programa": "simulado con las unidades del Día del VIN (ids ficticios)" if h.programa_simulado
            else "recibido", "unidades_programadas": h.programadas, "codigos_programados": len(h.filas),
            "cupo": h.cupo, "cupo_5": h.cupo_5,
            "sugeridas_total": sum(f.sugerida for f in h.filas) + len(h.exploracion),
            "codigos_con_cantidad": sum(1 for f in h.filas if f.sugerida),
            "filas_minimo_por_codigo": len(h.exploracion), "sin_cubrir": h.sin_cubrir,
            "minimo_por_codigo": {"P": h.minimo[0], "origen": h.minimo[1]}, "predictor": h.predictor,
            "ventana_tasa": list(h.ventana), "ventana_mercado": list(h.ventana_mercado),
            "archivos": [Path(a).name for a in archivos], "vin_en_salidas": 0,
            "semillas": {"programa": SEMILLA_PROGRAMA, "desempate": SEMILLA_DESEMPATE},
            "fuente_sha256": tabla.fuente.get("csv_sha256"), "catalogo_sha256": tabla.fuente.get("catalogo_sha256"),
            "version_codigo": version_codigo()}


def correr(tabla, opciones):
    """Demo de E3 para `opciones.dia_hoja` (190 por defecto) en `opciones.salida`; devuelve el resumen de p8.json."""
    dia = getattr(opciones, "dia_hoja", None) or DIA_DEMO
    h, archivos = construir_hoja(tabla, dia, opciones.salida, getattr(opciones, "programa", None),
                                 getattr(opciones, "cupo", None))
    return resumen(h, archivos, tabla)


def prueba(tabla, config):
    """Hook de P7: la E3 final con la ganadora preregistrada, el último día de la prueba <= 260 y su cifra.

    `config`: sección e3 del preregistro con "dia", "evaluacion" (la ganadora en el tramo <= 260) y "salida"
    (carpeta fuera del repo; por defecto la caché). Solo con la tabla desbloqueada.
    """
    from .puntaje import crear
    assert tabla.desbloqueada, "La E3 final solo se arma con el preregistro acordado"
    preregistro = config["preregistro"]
    g = preregistro["ganadora_en_prueba"]
    p = (preregistro.get("piezas", {}).get("p5") or {}).get("minimo_por_codigo", {}).get("P")
    minimo = (int(p), "preregistrado (P5)") if isinstance(p, (int, float)) else None
    carpeta = Path(config.get("salida") or Path.home() / ".cache" / "ford-predictive-quality" / "e3-final")
    assert datos.RAIZ not in carpeta.resolve().parents, "La hoja tiene tasas por código: va fuera del repo"
    h, archivos = construir_hoja(tabla, config["dia"], carpeta, predictor=crear(g["familia"], g["parametros"], tabla),
                                 minimo=minimo, evaluacion=config.get("evaluacion"))
    return {**resumen(h, archivos, tabla), "carpeta": str(carpeta)}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--csv", required=True, type=Path)
    parser.add_argument("--catalogo", required=True, type=Path)
    parser.add_argument("--dia", type=int, default=DIA_DEMO, help="Día del programa (la tasa usa resultados <= día−5)")
    parser.add_argument("--programa", type=Path, help="CSV con columnas unidad,codigo; sin él, se simula")
    parser.add_argument("--cupo", type=int, help="Cupo del día; sin él, el 5 %% del programa")
    parser.add_argument("--salida", required=True, type=Path, help="Carpeta fuera del repo (tiene tasas por código)")
    parser.add_argument("--cache", type=Path, default=Path.home() / ".cache" / "ford-predictive-quality")
    a = parser.parse_args(argv)
    assert datos.RAIZ not in a.salida.resolve().parents, "La hoja tiene tasas por código: va fuera del repo"
    tabla = datos.cargar(a.csv, a.catalogo, cache=a.cache)
    h, archivos = construir_hoja(tabla, a.dia, a.salida, a.programa, a.cupo)
    json.dump(resumen(h, archivos, tabla), sys.stdout, ensure_ascii=False, indent=2)
    print()


if __name__ == "__main__":
    main()
